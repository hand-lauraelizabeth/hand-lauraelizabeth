"""Offline synthetic CSV-to-XLSX interoperability checks for the Public Records Workspace.

Uses an independent CSV parser and openpyxl save/reopen/resave. This is not
native Microsoft Excel or Google Sheets application verification.
"""
import csv
import io
import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
FIXTURE = {
    'version': 1,
    'requests': [
        {'id': 'REQ-00001', 'jurisdiction': 'federal', 'agency': 'Example Office',
         'stage': 'submitted', 'submittedOn': '2026-10-20', 'agencyTrackingId': '000019',
         'description': 'Records, "as filed"\nSecond line'},
        {'id': 'REQ-00A2', 'jurisdiction': 'local', 'agency': 'Café Records Unit',
         'stage': 'draft', 'submittedOn': '', 'agencyTrackingId': '',
         'description': 'Budget notes — no decision implied'},
    ],
    'events': [
        {'id': 'EVT-0002', 'requestId': 'REQ-00001', 'date': '2026-10-21',
         'kind': 'note', 'sourceId': 'SRC-00007', 'description': 'Received, then reviewed'},
        {'id': 'EVT-0001', 'requestId': 'REQ-00001', 'date': '2026-10-20',
         'kind': 'submitted', 'sourceId': 'SRC-00008', 'description': 'Request submitted'},
    ],
}

def run_core(mode, payload):
    js = r"""
import { exportCSV, exportJSON } from './free-tools/public-records/core.mjs';
const input = JSON.parse(process.argv[1]);
if (process.argv[2] === 'positive') {
  process.stdout.write(JSON.stringify({ csv: exportCSV(input), json: exportJSON(input) }));
} else if (process.argv[2] === 'json') {
  process.stdout.write(JSON.stringify({ json: exportJSON(input) }));
} else {
  try { exportCSV(input); process.stdout.write(JSON.stringify({ blocked: false })); }
  catch (error) { process.stdout.write(JSON.stringify({ blocked: true, message: error.message })); }
}
"""
    result = subprocess.run(
        ['node', '--input-type=module', '-e', js, json.dumps(payload, ensure_ascii=False), mode],
        cwd=ROOT, text=True, capture_output=True, check=True, timeout=15)
    return json.loads(result.stdout)

def workbook_rows(path):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook.active
        assert all(cell.data_type != 'f' for row in sheet for cell in row), 'Unexpected formula cell'
        return [[cell.value if cell.value is not None else '' for cell in row] for row in sheet]
    finally:
        workbook.close()

class CSVInteroperability(unittest.TestCase):
    def test_csv_parser_and_xlsx_reopen_resave_preserve_records(self):
        exported = run_core('positive', FIXTURE)
        self.assertEqual(json.loads(exported['json']), FIXTURE)
        rows = list(csv.reader(io.StringIO(exported['csv'], newline=''), strict=True))
        self.assertEqual(len(rows), 5)
        self.assertTrue(all(len(row) == 10 for row in rows))
        self.assertEqual(rows[0], ['recordType', 'requestId', 'eventId', 'date',
                                   'jurisdiction', 'agency', 'stage', 'kind',
                                   'sourceId', 'description'])
        self.assertEqual([row[1] for row in rows[1:3]], ['REQ-00001', 'REQ-00A2'])
        self.assertEqual([row[2] for row in rows[3:]], ['EVT-0001', 'EVT-0002'])
        self.assertEqual([row[8] for row in rows[3:]], ['SRC-00008', 'SRC-00007'])
        self.assertEqual(rows[1][9], 'Records, "as filed"\nSecond line')
        self.assertEqual(rows[1][8], '000019')
        self.assertEqual(rows[2][5], 'Café Records Unit')
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / 'first.xlsx'
            second = Path(directory) / 'resaved.xlsx'
            book = Workbook()
            for row in rows:
                book.active.append(row)
            book.save(first)
            book.close()
            self.assertEqual(workbook_rows(first), rows)
            reopened = load_workbook(first, data_only=False)
            reopened.save(second)
            reopened.close()
            self.assertEqual(workbook_rows(second), rows)

    def test_spreadsheet_formula_payloads_fail_closed(self):
        unsafe = ['=2+2', '+SUM(1,2)', '-1+2', '@SUM(1,2)',
                  '\t=HYPERLINK("https://example.invalid")',
                  'safe,=2+2', 'safe\n=2+2', '＝2+2', '\u202e=2+2']
        for value in unsafe:
            with self.subTest(value=value):
                payload = json.loads(json.dumps(FIXTURE))
                payload['events'][0]['description'] = value
                result = run_core('negative', payload)
                self.assertTrue(result['blocked'], repr(value))
                self.assertIn('CSV export blocked', result['message'])

    def test_formula_risks_in_identifiers_and_metadata_are_blocked_without_json_loss(self):
        # CSV export must check every user-controlled column, not just descriptions.
        cases = [
            ('request ID', 'requests', 0, 'id', '=2+2'),
            ('event ID', 'events', 0, 'id', '+SUM(1,2)'),
            ('agency', 'requests', 0, 'agency', '＝2+2'),
            ('agency tracking ID', 'requests', 0, 'agencyTrackingId', '@SUM(1,2)'),
            ('request description', 'requests', 0, 'description', 'note\n=2+2'),
            ('event source ID', 'events', 0, 'sourceId', '\t=2+2'),
            ('event description', 'events', 0, 'description', 'note,=2+2'),
        ]
        for label, collection, index, field, unsafe in cases:
            with self.subTest(field=label):
                payload = json.loads(json.dumps(FIXTURE))
                payload[collection][index][field] = unsafe
                if label == 'request ID':
                    for event in payload['events']:
                        if event['requestId'] == 'REQ-00001':
                            event['requestId'] = unsafe
                blocked = run_core('negative', payload)
                self.assertTrue(blocked['blocked'], label)
                self.assertIn('CSV export blocked', blocked['message'])
                preserved = run_core('json', payload)
                self.assertEqual(json.loads(preserved['json']), payload, label)

if __name__ == '__main__':
    test_count = unittest.defaultTestLoader.loadTestsFromTestCase(CSVInteroperability).countTestCases()
    if test_count < 3:
        raise SystemExit(f'Expected at least 3 interoperability tests, found {test_count}')
    unittest.main(verbosity=2)
